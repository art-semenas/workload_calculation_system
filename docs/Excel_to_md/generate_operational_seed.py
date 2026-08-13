#!/usr/bin/env python3
"""Regenerate the PoC operational-data seed from the source workbook.

    pip install openpyxl
    python docs/Excel_to_md/generate_operational_seed.py

Reads  docs/Шаблон нагрузки_з_v_4_00.xlsx  (sheets Дорога, ОС, ПС, Видео, Ремонт, Записи)
Writes backend/src/main/resources/db/changelog/changes/v1.0.9-seed-operational-data.sql

Covers import steps 6-9 of the extraction spec: travel, device quantities, repair counts and
records counts. Joins to the objects seeded by v1.0.6 through `import_seq_no`, and to the
catalogues seeded by v1.0.1 by name.

The output is a generated artefact — edit this script, not the SQL. See
docs/Excel_to_md/Шаблон_нагрузки_v4_data_extraction_spec.md §4, §7 and §9.
"""
import openpyxl, io, re, warnings, collections, os

warnings.filterwarnings("ignore")

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, "docs", "Шаблон нагрузки_з_v_4_00.xlsx")
OUT = os.path.join(ROOT, "backend", "src", "main", "resources", "db", "changelog",
                   "changes", "v1.0.9-seed-operational-data.sql")

LAST = 2935  # last spreadsheet row; data starts at row 2

wb = openpyxl.load_workbook(PATH, data_only=True)


def num(v):
    """Numeric value, or None for blanks and the status text described in spec §7.1."""
    return v if isinstance(v, (int, float)) else None


def q(s):
    return "'" + str(s).replace("'", "''") + "'"


# --------------------------------------------------------------------------------------
# Column header -> seeded catalogue name.
# The workbook headers carry a ", ед" unit suffix and, for repairs, longer wording than the
# names seeded in v1.0.1. Mapping explicitly rather than by fuzzy match, because a silent
# mismatch would drop rows.
# --------------------------------------------------------------------------------------
DEVICE_SHEETS = {
    # sheet: (system_type, first col, last col)
    "ОС": ("OS", 6, 19),
    "ПС": ("PS", 6, 23),
    "Видео": ("VIDEO", 6, 8),
}

REPAIR_NAME_MAP = {
    "Замена приемно-контрольного прибора серии А6 ОС": "Замена ПКП серии А6 ОС",
    "Замена приемно-контрольного прибора серии А6 ПС": "Замена ПКП серии А6 ПС",
    "Замена шунтирующих и оконечного резисторов шлейфа сигнализации ОС":
        "Замена шунтирующих/оконечного резисторов шлейфа ОС",
    "Замена шунтирующих и оконечного резисторов шлейфа сигнализации ПС":
        "Замена шунтирующих/оконечного резисторов шлейфа ПС",
    "Замена основных составных частей видеосервера в комплекте":
        "Замена основных составных частей видеосервера в комплексе",
    "Переустановка программного обеспечения на видеосервере": "Переустановка ПО на видеосервере",
    "Дефектный акт (с учетом установления причин неисправности) ОС": "Дефектный акт ОС",
    "Дефектный акт (с учетом установления причин неисправности) ПС": "Дефектный акт ПС",
    "Дефектный акт (с учетом установления причин неисправности) Видео": "Дефектный акт Видео",
}

RECORDS_COLUMNS = [
    (6, "access_requests"),
    (7, "monitoring_requests"),
    (8, "footage_requests"),
    (9, "backup_control"),
    (10, "security_admin"),
]


def header(ws, col):
    return " ".join(str(ws.cell(row=1, column=col).value).replace("­", "").split())


def catalog_name(raw):
    """Strip the workbook's unit suffix to reach the seeded catalogue name."""
    return re.sub(r",\s*ед\d*$", "", raw).strip()


def seq_of(ws, r):
    return ws.cell(row=r, column=1).value


# --------------------------------------------------------------------------------------
# Collect
# --------------------------------------------------------------------------------------
stats = collections.Counter()

# --- travel -------------------------------------------------------------------------
dor = wb["Дорога"]
travel = []
for r in range(2, LAST + 1):
    dist = num(dor.cell(row=r, column=7).value)
    oneway = num(dor.cell(row=r, column=8).value)
    if dist is None and oneway is None:
        stats["travel_skipped"] += 1
        continue
    travel.append((seq_of(dor, r), dist, oneway))
stats["travel"] = len(travel)

# --- device quantities --------------------------------------------------------------
# (seq, system, catalogue_name, qty). One row per non-zero cell.
assignments = []
for sheet, (system, c1, c2) in DEVICE_SHEETS.items():
    ws = wb[sheet]
    for c in range(c1, c2 + 1):
        name = catalog_name(header(ws, c))
        for r in range(2, LAST + 1):
            v = num(ws.cell(row=r, column=c).value)
            if v is None:
                if ws.cell(row=r, column=c).value is not None:
                    stats["device_text_cells"] += 1
                continue
            if v <= 0:
                continue
            assignments.append((seq_of(ws, r), system, name, v))
stats["assignments"] = len(assignments)

# object_devices holds one physical count per (object, device_type). A device assigned to
# two systems keeps the first occurrence's quantity (TOR §11.1 upsert rule).
physical = {}
for seq, system, name, qty in assignments:
    physical.setdefault((seq, name), qty)
stats["object_devices"] = len(physical)

# --- repair counts ------------------------------------------------------------------
rem = wb["Ремонт"]
repairs = []
for c in range(6, 31):  # F:AD — 16 work types then 9 document types
    raw = header(rem, c)
    name = REPAIR_NAME_MAP.get(raw, raw)
    for r in range(2, LAST + 1):
        v = num(rem.cell(row=r, column=c).value)
        if v is None:
            if rem.cell(row=r, column=c).value is not None:
                stats["repair_text_cells"] += 1
            continue
        if v <= 0:
            continue
        repairs.append((seq_of(rem, r), name, int(v)))
stats["repairs"] = len(repairs)

# --- records counts -----------------------------------------------------------------
zap = wb["Записи"]
records = []
for r in range(2, LAST + 1):
    vals = []
    any_positive = False
    for c, _ in RECORDS_COLUMNS:
        v = num(zap.cell(row=r, column=c).value)
        if v is None and zap.cell(row=r, column=c).value is not None:
            stats["records_text_cells"] += 1
        v = v or 0
        if v > 0:
            any_positive = True
        vals.append(v)
    if any_positive:
        records.append((seq_of(zap, r), vals))
stats["records"] = len(records)

# --------------------------------------------------------------------------------------
# Emit
# --------------------------------------------------------------------------------------
L = []
w = L.append
w("--liquibase formatted sql")
w("")
w("-- PoC operational-data seed, extracted from docs/Шаблон нагрузки_з_v_4_00.xlsx.")
w("-- Generated artefact — regenerate with docs/Excel_to_md/generate_operational_seed.py.")
w("-- Objects join by import_seq_no (v1.0.6); catalogues join by name (v1.0.1).")
w("-- Status text in quantity cells is skipped, never coerced to a number — see spec §7.1.")
w("")

w("--changeset a.semenas:v1.0.9-1 context:demo-data")
w("--comment: Seed travel distance and one-way time (%d objects)" % len(travel))
w("INSERT INTO travel (object_id, distance_km, one_way_time_min)")
w("SELECT o.id, v.distance_km, v.one_way_min")
w("FROM (VALUES")
w(",\n".join("  (%d, %s, %s)" % (s, "NULL" if d is None else repr(float(d)),
                                 "NULL" if t is None else repr(float(t)))
             for s, d, t in travel))
w(") AS v(seq, distance_km, one_way_min)")
w("JOIN objects o ON o.import_seq_no = v.seq;")
w("--rollback DELETE FROM travel;")
w("")

w("--changeset a.semenas:v1.0.9-2 context:demo-data")
w("--comment: Seed physical device counts (%d object-device pairs)" % len(physical))
w("INSERT INTO object_devices (object_id, device_type_id, quantity_physical)")
w("SELECT o.id, dt.id, v.qty")
w("FROM (VALUES")
w(",\n".join("  (%d, %s, %s)" % (seq, q(name), float(qty))
             for (seq, name), qty in sorted(physical.items())))
w(") AS v(seq, device_name, qty)")
w("JOIN objects o      ON o.import_seq_no = v.seq")
w("JOIN device_types dt ON dt.name = v.device_name;")
w("--rollback DELETE FROM object_devices;")
w("")

w("--changeset a.semenas:v1.0.9-3 context:demo-data")
w("--comment: Seed maintained quantities per system (%d assignments)" % len(assignments))
w("INSERT INTO object_system_assignments (object_id, device_type_id, system_type, quantity_maintained, context_id)")
w("SELECT o.id, dt.id, v.system_type, v.qty, dsc.id")
w("FROM (VALUES")
w(",\n".join("  (%d, %s, %s, %s)" % (seq, q(system), q(name), float(qty))
             for seq, system, name, qty in assignments))
w(") AS v(seq, system_type, device_name, qty)")
w("JOIN objects o                ON o.import_seq_no = v.seq")
w("JOIN device_types dt          ON dt.name = v.device_name")
w("JOIN device_system_contexts dsc ON dsc.device_type_id = dt.id AND dsc.system_type = v.system_type;")
w("--rollback DELETE FROM object_system_assignments;")
w("")

w("--changeset a.semenas:v1.0.9-4 context:demo-data")
w("--comment: Seed repair counts for the period (%d rows)" % len(repairs))
w("INSERT INTO object_repairs (object_id, repair_type_id, count)")
w("SELECT o.id, rt.id, v.cnt")
w("FROM (VALUES")
w(",\n".join("  (%d, %s, %d)" % (seq, q(name), cnt) for seq, name, cnt in repairs))
w(") AS v(seq, repair_name, cnt)")
w("JOIN objects o       ON o.import_seq_no = v.seq")
w("JOIN repair_types rt ON rt.name = v.repair_name;")
w("--rollback DELETE FROM object_repairs;")
w("")

w("--changeset a.semenas:v1.0.9-5 context:demo-data")
w("--comment: Seed records/admin task counts (%d objects with activity)" % len(records))
w("INSERT INTO records_tasks (object_id, access_requests, monitoring_requests, footage_requests, backup_control, security_admin)")
w("SELECT o.id, v.access_requests, v.monitoring_requests, v.footage_requests, v.backup_control, v.security_admin")
w("FROM (VALUES")
w(",\n".join("  (%d, %s, %s, %s, %s, %s)" % (seq, *[float(x) for x in vals])
             for seq, vals in records))
w(") AS v(seq, access_requests, monitoring_requests, footage_requests, backup_control, security_admin)")
w("JOIN objects o ON o.import_seq_no = v.seq;")
w("--rollback DELETE FROM records_tasks;")
w("")

io.open(OUT, "w", encoding="utf-8", newline="\n").write("\n".join(L))

for k in ["travel", "travel_skipped", "object_devices", "assignments", "repairs", "records",
          "device_text_cells", "repair_text_cells", "records_text_cells"]:
    print("%-20s %d" % (k, stats[k]))
print("wrote %s" % os.path.relpath(OUT, ROOT))
