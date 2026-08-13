#!/usr/bin/env python3
"""Regenerate the PoC org-data seed from the source workbook.

    pip install openpyxl
    python docs/Excel_to_md/generate_org_seed.py

Reads  docs/Шаблон нагрузки_з_v_4_00.xlsx  (sheet ОС cols A:E, sheet Дорога for repairs)
Writes backend/src/main/resources/db/changelog/changes/v1.0.6-seed-org-data.sql

The output is a generated artefact — edit this script, not the SQL. See
docs/Excel_to_md/Шаблон_нагрузки_v4_data_extraction_spec.md §3 and §10 for the rules
this implements and §7 for the data-quality issues it works around.
"""
import openpyxl, io, re, warnings, collections, os, sys
warnings.filterwarnings("ignore")

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PATH = os.path.join(ROOT, "docs", "Шаблон нагрузки_з_v_4_00.xlsx")
OUT  = os.path.join(ROOT, "backend", "src", "main", "resources", "db", "changelog",
                    "changes", "v1.0.6-seed-org-data.sql")

wb = openpyxl.load_workbook(PATH, data_only=True)
LAST = 2935
oc, dor = wb["ОС"], wb["Дорога"]

def clean(v):
    if v is None: return ""
    return " ".join(str(v).replace("\u00ad", "").split())

rows = []
for r in range(2, LAST + 1):
    div = clean(oc.cell(row=r, column=2).value)
    nm  = clean(oc.cell(row=r, column=4).value)
    # repair the two formula-paste corruptions from the Дорога sheet (spec §3.1)
    if "B1830" in div: div = clean(dor.cell(row=r, column=2).value)
    if "D1834" in nm:  nm  = clean(dor.cell(row=r, column=4).value)
    rows.append(dict(
        seq=oc.cell(row=r, column=1).value,
        div=div,
        br=clean(oc.cell(row=r, column=3).value),
        raw=nm,
        eng=clean(oc.cell(row=r, column=5).value),
    ))

# ---------------- name / address split ----------------
PAT_OTD  = re.compile(r'^Отделение\s*№?\s*(\d+)\s*/\s*(\d+)\s*(.*)$')
PAT_NUM  = re.compile(r'^(\d{3,6})\s+(.*)$')
GEO      = re.compile(r'\s*(?=(?:\bг\.|\bаг\.|\bд\.\s|\bгп\.|\bг\.п\.|\bп\.\s|\bкп\.|\bпгт|\bМинск\b|\bБарановичи\b))')

def split_name(s):
    """name keeps the source label verbatim; address is a derived best-effort extract.

    The workbook packs facility code, label and address into one free-text cell with no
    consistent separator (1332 rows lead with an internal facility code, 1054 with
    'Отделение №X/Y', the rest with a descriptor). Any split that produces a short name
    discards either the code or the descriptor, so `name` stays whole — nothing is lost —
    and `address` carries the geographic substring for search and display.
    """
    s = s.strip().rstrip('*').strip().rstrip(',').strip()
    for pat in (PAT_OTD, PAT_NUM):
        m = pat.match(s)
        if m:
            return s, m.group(3 if pat is PAT_OTD else 2).strip(" ,.")
    parts = GEO.split(s, 1)
    if len(parts) == 2 and parts[0].strip():
        return s, parts[1].strip(" ,.")
    return s, ""

for x in rows:
    x["name"], x["address"] = split_name(x["raw"])

# ---------------- engineers ----------------
# Source typo: one row drops the leading 'А' of 'Александр Н Соловей' (1 object vs 101).
# Same person — merge before building the account list.
ENGINEER_FIXES = {
    "лександр Н Соловей/BELARUSBANK/BY": "Александр Н Соловей/BELARUSBANK/BY",
}
for x in rows:
    if x["eng"] in ENGINEER_FIXES:
        x["eng"] = ENGINEER_FIXES[x["eng"]]

TRANSLIT = {'а':'a','б':'b','в':'v','г':'g','д':'d','е':'e','ё':'e','ж':'zh','з':'z','и':'i','й':'i',
            'к':'k','л':'l','м':'m','н':'n','о':'o','п':'p','р':'r','с':'s','т':'t','у':'u','ф':'f',
            'х':'kh','ц':'ts','ч':'ch','ш':'sh','щ':'shch','ъ':'','ы':'y','ь':'','э':'e','ю':'yu','я':'ya'}
def translit(s):
    return "".join(TRANSLIT.get(ch, ch) for ch in s.lower())
def slug(s):
    # 'Сергей_Ца А Широков' carries a data-entry artefact in the given name; the display
    # name is preserved verbatim, but the login takes only the part before the underscore.
    return re.sub(r'[^a-z0-9]+', '', translit(s.split('_')[0]))

# home division = division where the engineer services the most objects
eng_div = collections.defaultdict(collections.Counter)
for x in rows:
    if x["eng"]: eng_div[x["eng"]][x["div"]] += 1

engineers, used = {}, {}
for src in sorted(eng_div, key=lambda e: (-sum(eng_div[e].values()), e)):
    core = src.replace("/BELARUSBANK/BY", "").strip()
    toks = core.split()
    first, last = toks[0], toks[-1]
    base = "%s.%s" % (slug(first), slug(last))
    email = base + "@workload.local"
    if email in used:                       # 1 known collision: Сергей * Карпук
        mid = slug(toks[1]) if len(toks) == 3 else ""
        email = "%s.%s.%s@workload.local" % (slug(first), mid, slug(last))
    used[email] = src
    engineers[src] = dict(email=email, name=core,
                          home=eng_div[src].most_common(1)[0][0],
                          employee_id=src)

# ---------------- SQL emit ----------------
def q(s):
    return "'" + str(s).replace("'", "''") + "'"

divisions = sorted({x["div"] for x in rows})
branches  = sorted({(x["div"], x["br"]) for x in rows})

L = []
w = L.append
w("--liquibase formatted sql")
w("")
w("-- PoC org-data seed, extracted from docs/Шаблон нагрузки_з_v_4_00.xlsx (sheet ОС, cols A:E).")
w("-- Generated artefact — see docs/Excel_to_md/Шаблон_нагрузки_v4_data_extraction_spec.md §3.")
w("-- Foreign keys resolve by name, so this file is independent of generated UUIDs.")
w("")

w("--changeset a.semenas:v1.0.6-1 context:demo-data")
w("--comment: Seed divisions (%d)" % len(divisions))
w("INSERT INTO divisions (name) VALUES")
w(",\n".join("  (%s)" % q(d) for d in divisions) + ";")
w("--rollback DELETE FROM divisions;")
w("")

w("--changeset a.semenas:v1.0.6-2 context:demo-data")
w("--comment: Seed branches (%d) — all 135 names are globally unique, but the FK is resolved by (division, name) anyway" % len(branches))
w("INSERT INTO branches (division_id, name)")
w("SELECT d.id, v.branch_name")
w("FROM (VALUES")
w(",\n".join("  (%s, %s)" % (q(dv), q(br)) for dv, br in branches))
w(") AS v(division_name, branch_name)")
w("JOIN divisions d ON d.name = v.division_name;")
w("--rollback DELETE FROM branches;")
w("")

w("--changeset a.semenas:v1.0.6-3 context:demo-data")
w("--comment: Seed engineer placeholder accounts (%d) — TOR C-31: requires_activation, unusable password" % len(engineers))
w("INSERT INTO users (email, name, password_hash, role, home_division_id, capacity_fte, employee_id, is_active, requires_activation)")
w("SELECT v.email, v.full_name, '!', 'engineer', d.id, 1.00, v.employee_id, TRUE, TRUE")
w("FROM (VALUES")
w(",\n".join("  (%s, %s, %s, %s)" % (q(e["email"]), q(e["name"]), q(e["employee_id"]), q(e["home"]))
             for e in sorted(engineers.values(), key=lambda x: x["email"])))
w(") AS v(email, full_name, employee_id, home_division)")
w("JOIN divisions d ON d.name = v.home_division;")
w("--rollback DELETE FROM users WHERE role = 'engineer' AND email LIKE '%@workload.local';")
w("")

w("--changeset a.semenas:v1.0.6-4 context:demo-data")
w("--comment: Seed objects (%d) — import_seq_no preserves the workbook row order (№)" % len(rows))
w("INSERT INTO objects (branch_id, name, address, import_seq_no)")
w("SELECT b.id, v.name, NULLIF(v.address, ''), v.seq")
w("FROM (VALUES")
w(",\n".join("  (%d, %s, %s, %s, %s)" % (x["seq"], q(x["div"]), q(x["br"]), q(x["name"]), q(x["address"]))
             for x in rows))
w(") AS v(seq, division_name, branch_name, name, address)")
w("JOIN divisions d ON d.name = v.division_name")
w("JOIN branches  b ON b.division_id = d.id AND b.name = v.branch_name;")
w("--rollback DELETE FROM objects;")
w("")

assigned = [x for x in rows if x["eng"]]
w("--changeset a.semenas:v1.0.6-5 context:demo-data")
w("--comment: Seed object-engineer assignments (%d of %d objects, %d unassigned in source)"
  % (len(assigned), len(rows), len(rows) - len(assigned)))
w("INSERT INTO object_engineers (object_id, engineer_id)")
w("SELECT o.id, u.id")
w("FROM (VALUES")
w(",\n".join("  (%d, %s)" % (x["seq"], q(engineers[x["eng"]]["email"])) for x in assigned))
w(") AS v(seq, email)")
w("JOIN objects o ON o.import_seq_no = v.seq")
w("JOIN users   u ON u.email = v.email;")
w("--rollback DELETE FROM object_engineers;")
w("")

io.open(OUT, "w", encoding="utf-8", newline="\n").write("\n".join(L))

# ---------------- summary ----------------
print("divisions   %d" % len(divisions))
print("branches    %d" % len(branches))
print("engineers   %d" % len(engineers))
print("objects     %d" % len(rows))
print("assignments %d (unassigned %d)" % (len(assigned), len(rows) - len(assigned)))
print("address extracted: %d / %d" % (sum(1 for x in rows if x["address"]), len(rows)))
print("emails unique: %s" % (len({e["email"] for e in engineers.values()}) == len(engineers)))
print("wrote %s" % os.path.relpath(OUT, ROOT))
